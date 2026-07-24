from __future__ import annotations

import re
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"C:\Users\lenovo\Documents\Obsidian Vault\20_Projects\当前项目\小程序功能与业务流程说明（客户版）.md")
OUTPUT = Path(r"C:\Users\lenovo\Documents\Education System\docs\deliverables\小程序功能与业务流程说明（客户版）.docx")
ASSET_DIR = Path(r"C:\Users\lenovo\AppData\Local\Temp\EducationSystemClientGuide")

# Preset: compact_reference_guide.
# Named overrides: Microsoft YaHei for East Asian glyphs, restrained project green
# for the first-page kicker/callouts, and customer_pack first-page furniture.
COLORS = {
    "navy": "0B2545",
    "blue": "2E74B5",
    "dark_blue": "1F4D78",
    "green": "159A62",
    "green_dark": "0E7650",
    "green_light": "EAF7F0",
    "blue_light": "E8EEF5",
    "orange": "F28C1B",
    "gold_light": "FFF7E7",
    "ink": "243342",
    "muted": "667085",
    "line": "D8E2EA",
    "white": "FFFFFF",
}


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, *, size=None, color=None, bold=None, italic=None, east_asia="Microsoft YaHei"):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_shading(paragraph, fill: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D8E2EA", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, (cell, width) in enumerate(zip(row.cells, widths_dxa)):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=COLORS["muted"])
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr_text, fld_sep, fld_text, fld_end])
    end_run = paragraph.add_run(" 页")
    set_run_font(end_run, size=9, color=COLORS["muted"])


def style_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(COLORS["ink"])
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, COLORS["blue"], 18, 10),
        "Heading 2": (13, COLORS["blue"], 14, 7),
        "Heading 3": (12, COLORS["dark_blue"], 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    run = hp.add_run("21天自主学习训练营  |  功能与业务流程说明")
    set_run_font(run, size=9, color=COLORS["muted"], bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.paragraph_format.space_before = Pt(0)
    add_page_field(fp)


def add_inline_markdown(paragraph, text: str, *, base_size=11, base_color=None, bold=False):
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        is_bold = part.startswith("**") and part.endswith("**")
        is_code = part.startswith("`") and part.endswith("`")
        clean = part[2:-2] if is_bold else (part[1:-1] if is_code else part)
        run = paragraph.add_run(clean)
        set_run_font(
            run,
            size=base_size,
            color=base_color or COLORS["ink"],
            bold=bold or is_bold,
            east_asia="Microsoft YaHei",
        )
        if is_code:
            run.font.name = "Consolas"
            run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Consolas")
            run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")


def add_callout(doc, title: str, body: str, fill="EAF7F0", accent="159A62"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360], indent_dxa=120)
    set_table_borders(table, color=accent, size="8")
    cell = table.cell(0, 0)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, bottom=150, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    add_inline_markdown(p, title, base_size=11, base_color=accent, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.2
    add_inline_markdown(p2, body, base_size=10.5, base_color=COLORS["ink"])
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_customer_pack_title(doc: Document):
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(8)
    kicker.paragraph_format.space_after = Pt(2)
    run = kicker.add_run("客户功能说明")
    set_run_font(run, size=11, color=COLORS["green"], bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run("小程序功能与业务流程说明")
    set_run_font(run, size=28, color=COLORS["navy"], bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    run = subtitle.add_run("21天自主学习训练营 · 用户端、支付端、会员端与运营端全流程")
    set_run_font(run, size=13, color=COLORS["muted"])

    table = doc.add_table(rows=2, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [1300, 3380, 1300, 3380], indent_dxa=0)
    set_table_borders(table, color=COLORS["line"], size="5")
    metadata = [
        ("文档用途", "客户了解与项目验收", "版本", "V1.3"),
        ("适用范围", "微信小程序及运营后台", "日期", "2026年7月23日"),
    ]
    for row, values in zip(table.rows, metadata):
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, bottom=100, start=120, end=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            is_label = idx % 2 == 0
            if is_label:
                set_cell_shading(cell, COLORS["blue_light"])
            add_inline_markdown(
                p,
                value,
                base_size=9.5 if is_label else 10,
                base_color=COLORS["dark_blue"] if is_label else COLORS["ink"],
                bold=is_label,
            )

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    add_callout(
        doc,
        "一句话说明",
        "用户先体验内容与测评，再报名训练营；支付成功后系统自动开通永久会员并进入服务流程，会员还可通过专属海报邀请好友并按规则获得奖励。",
    )


def load_font(size: int, bold=False):
    candidates = [
        Path(r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def draw_arrow(draw, start, end, color, width=6):
    draw.line([start, end], fill=color, width=width)
    x2, y2 = end
    x1, y1 = start
    if abs(x2 - x1) >= abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        points = [(x2, y2), (x2 - 18 * direction, y2 - 12), (x2 - 18 * direction, y2 + 12)]
    else:
        direction = 1 if y2 > y1 else -1
        points = [(x2, y2), (x2 - 12, y2 - 18 * direction), (x2 + 12, y2 - 18 * direction)]
    draw.polygon(points, fill=color)


def create_flow_diagram(path: Path, title: str, steps, colors):
    width, height = 1600, 420
    image = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    title_font = load_font(34, bold=True)
    step_font = load_font(26, bold=True)
    caption_font = load_font(18, bold=False)
    draw.text((60, 34), title, fill="#0B2545", font=title_font)

    margin_x = 60
    gap = 24
    box_y = 135
    box_h = 170
    box_w = int((width - margin_x * 2 - gap * (len(steps) - 1)) / len(steps))

    for idx, step in enumerate(steps):
        x = margin_x + idx * (box_w + gap)
        fill = colors[idx % len(colors)]
        draw.rounded_rectangle((x, box_y, x + box_w, box_y + box_h), radius=24, fill=fill, outline="#D8E2EA", width=2)
        number_fill = "#159A62" if idx % 3 == 0 else ("#2D80E8" if idx % 3 == 1 else "#F28C1B")
        draw.ellipse((x + 18, box_y + 18, x + 66, box_y + 66), fill=number_fill)
        num_text = str(idx + 1)
        bbox = draw.textbbox((0, 0), num_text, font=caption_font)
        draw.text((x + 42 - (bbox[2] - bbox[0]) / 2, box_y + 42 - (bbox[3] - bbox[1]) / 2 - 2), num_text, fill="white", font=caption_font)

        lines = step if isinstance(step, (list, tuple)) else [step]
        total_h = len(lines) * 38
        current_y = box_y + 92 - total_h / 2
        for line in lines:
            bbox = draw.textbbox((0, 0), line, font=step_font)
            text_x = x + box_w / 2 - (bbox[2] - bbox[0]) / 2
            draw.text((text_x, current_y), line, fill="#243342", font=step_font)
            current_y += 38

        if idx < len(steps) - 1:
            draw_arrow(draw, (x + box_w + 3, box_y + box_h / 2), (x + box_w + gap - 3, box_y + box_h / 2), "#96A3AF", 5)

    image.save(path, quality=95)


def build_diagrams():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    diagrams = [
        ASSET_DIR / "workflow-overview.png",
        ASSET_DIR / "workflow-payment.png",
        ASSET_DIR / "workflow-referral.png",
    ]
    create_flow_diagram(
        diagrams[0],
        "小程序四阶段服务闭环",
        [["内容体验", "测评工具"], ["训练营", "报名支付"], ["会员服务", "登记开营"], ["邀请推广", "收益提现"]],
        ["#EAF7F0", "#EDF4FE", "#FFF7E7", "#F2F4F7"],
    )
    create_flow_diagram(
        diagrams[1],
        "报名支付与会员开通链路",
        [["选择永久", "会员方案"], ["创建训练营", "会员订单"], ["微信", "虚拟支付"], ["服务端", "确认到账"], ["自动开通", "永久会员"], ["填写", "报名登记"]],
        ["#EAF7F0", "#EDF4FE", "#FFF7E7"],
    )
    create_flow_diagram(
        diagrams[2],
        "会员邀请与奖励链路",
        [["会员生成", "专属海报"], ["二维码携带", "推荐参数"], ["好友扫码", "进入小程序"], ["系统记录", "推荐关系"], ["好友完成", "有效报名"], ["奖励结算", "申请提现"]],
        ["#EAF7F0", "#EDF4FE", "#FFF7E7"],
    )
    return diagrams


def add_diagram(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.45))
    doc_pr = run._r.xpath(".//wp:docPr")
    if doc_pr:
        doc_pr[0].set("descr", caption)
    p.paragraph_format.space_after = Pt(8)


def markdown_body_lines():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    if lines and lines[0].strip() == "---":
        end = next(i for i in range(1, len(lines)) if lines[i].strip() == "---")
        lines = lines[end + 1 :]
    while lines and not lines[0].strip():
        lines.pop(0)
    if lines and lines[0].startswith("# "):
        lines.pop(0)
    return lines


def build_document():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    diagrams = build_diagrams()

    doc = Document()
    style_document(doc)
    add_customer_pack_title(doc)

    lines = markdown_body_lines()
    diagram_index = 0
    diagram_captions = [
        "图 1  小程序四阶段服务闭环",
        "图 2  报名支付与会员开通链路",
        "图 3  会员邀请与奖励链路",
    ]

    i = 0
    current_section = ""
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        if stripped.startswith("```mermaid"):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                i += 1
            if diagram_index < len(diagrams):
                add_diagram(doc, diagrams[diagram_index], diagram_captions[diagram_index])
                diagram_index += 1
            i += 1
            continue

        if stripped.startswith("## "):
            title = stripped[3:].strip()
            current_section = title
            p = doc.add_paragraph(style="Heading 1")
            if title.startswith(("十一、", "十二、")):
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(6)
            add_inline_markdown(p, title, base_size=16, base_color=COLORS["blue"], bold=True)
            i += 1
            continue

        if stripped.startswith("### "):
            title = stripped[4:].strip()
            p = doc.add_paragraph(style="Heading 2")
            add_inline_markdown(p, title, base_size=13, base_color=COLORS["blue"], bold=True)
            i += 1
            continue

        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_markdown(p, stripped[2:].strip())
            i += 1
            continue

        numbered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if numbered:
            # Keep the source document's explicit numbering. Word's built-in
            # List Number style continues counters across separate sections,
            # which can make a new checklist incorrectly start at 5 or 13.
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.125)
            compact_checklist = current_section.startswith(("十一、", "十二、"))
            p.paragraph_format.space_after = Pt(0 if compact_checklist else 2)
            p.paragraph_format.line_spacing = 1.0 if compact_checklist else 1.15
            list_size = 10.5 if compact_checklist else 11
            prefix = p.add_run(f"{numbered.group(1)}. ")
            set_run_font(prefix, size=list_size, color=COLORS["ink"])
            add_inline_markdown(p, numbered.group(2).strip(), base_size=list_size)
            i += 1
            continue

        paragraph_lines = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt.startswith(("## ", "### ", "- ", "```mermaid")) or re.match(r"^\d+\.\s+", nxt):
                break
            paragraph_lines.append(nxt)
            i += 1
        text = "".join(paragraph_lines)
        p = doc.add_paragraph()
        add_inline_markdown(p, text)

    # Keep document properties neutral and client-ready.
    props = doc.core_properties
    props.title = "小程序功能与业务流程说明（客户版）"
    props.subject = "21天自主学习训练营小程序功能与业务流程"
    props.author = "Education System 项目组"
    props.keywords = "微信小程序, 训练营, 会员, 支付, 邀请, 提现"
    props.comments = "客户功能说明 V1.3"

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    try:
        output = build_document()
        print(output)
    except Exception as exc:
        print(f"BUILD_FAILED: {exc}", file=sys.stderr)
        raise
